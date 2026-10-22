from __future__ import annotations

import os
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .metrics import clock_from_second, datetime_label, format_duration
from .models import DeviceResult, Segment, SiteConfig
from .rules import DAY


TITLE_FONT = Font(name="微软雅黑", bold=True, size=14)
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="微软雅黑", size=10)
NOTE_FONT = Font(name="微软雅黑", italic=True, size=9, color="666666")
WARN_FONT = Font(name="微软雅黑", size=9, color="C00000")

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
BAND_FILL = PatternFill("solid", fgColor="2F5597")
LIGHT_FILL = PatternFill("solid", fgColor="F2F2F2")

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def generate_outputs(
    results: tuple[DeviceResult, ...],
    config: SiteConfig,
    output_dir: Path,
    generated_at: datetime,
    force: bool = False,
) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / config.report_name
    progress_path = output_dir / config.progress_name
    existing = [path for path in (report_path, progress_path) if path.exists()]
    if existing and not force:
        names = ", ".join(path.name for path in existing)
        raise FileExistsError(f"Output already exists: {names}. Use --force to replace it.")

    report = _build_report_workbook(results, config, generated_at)
    progress = _build_progress_workbook(results, config, generated_at)
    _atomic_save(report, report_path)
    _atomic_save(progress, progress_path)
    return report_path, progress_path


def _atomic_save(workbook: Workbook, path: Path) -> None:
    temporary = path.with_name(f".{path.stem}.tmp{path.suffix}")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        if temporary.exists():
            temporary.unlink()


def _build_report_workbook(
    results: tuple[DeviceResult, ...],
    config: SiteConfig,
    generated_at: datetime,
) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for result in results:
        sheet = workbook.create_sheet(_safe_sheet_name(result.config.display_name))
        _write_device_sheet(sheet, result)
    workbook.properties.title = f"{config.site_name}设备日志分析报告"
    workbook.properties.creator = "device-log-analyzer"
    workbook.properties.created = generated_at
    workbook.properties.modified = generated_at
    return workbook


def _safe_sheet_name(name: str) -> str:
    for character in "\\/*?:[]":
        name = name.replace(character, "_")
    return name[:31] or "设备"


def _write_device_sheet(sheet: Worksheet, result: DeviceResult) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

    erasures = result.timeline.erasures
    window_count = len(erasures) + 1
    total_columns = 1 + window_count + 3
    if not erasures:
        total_columns += 1
    total_columns = max(total_columns, 6)

    origin = result.actual_days[0]
    start_date = result.config.expected_start or origin
    end_date = result.config.expected_end or result.actual_days[-1]
    title = (
        f"{result.config.display_name} 设备日志分析"
        if result.config.display_name == result.config.id
        else f"{result.config.display_name}  {result.config.id} 设备日志分析"
    )
    row = 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    cell = sheet.cell(row, 1, title)
    cell.font = TITLE_FONT
    cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[row].height = 26

    row += 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    sheet.cell(
        row,
        1,
        f"日志时间范围：{start_date.isoformat()} 00:00:00 ~ {end_date.isoformat()} 23:59:59",
    ).font = BODY_FONT

    row += 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    missing = (
        "无"
        if not result.missing_days
        else "、".join(day.strftime("%m-%d") for day in result.missing_days)
    )
    sheet.cell(
        row,
        1,
        f"缺失日期（共 {len(result.missing_days)} 天）：{missing}",
    ).font = BODY_FONT

    row += 2
    row = _section(sheet, row, "一、擦写统计", total_columns)
    first_value = result.timeline.first_value
    details = "；".join(
        f"第{index}次 {datetime_label(origin, change.global_second)} "
        f"avgEraseCount {change.before} → {change.after}"
        for index, change in enumerate(erasures, 1)
    )
    config_time = result.config_time_label or "未发现配置保存记录"
    first_after_config = (
        datetime_label(origin, result.first_erase_after_config)
        if result.first_erase_after_config is not None
        else "未发生"
    )
    _key_value(
        sheet,
        row,
        "起始 avgEraseCount",
        first_value if first_value is not None else "—",
        total_columns,
    )
    row += 1
    _key_value(sheet, row, "擦写次数", len(erasures), total_columns)
    row += 1
    _key_value(sheet, row, "擦写明细", details or "本阶段无擦写", total_columns)
    row += 1
    _key_value(sheet, row, "配置更改完成时间", config_time, total_columns)
    row += 1
    _key_value(
        sheet,
        row,
        "配置后第一次擦写时间",
        first_after_config,
        total_columns,
    )

    row += 2
    row = _section(sheet, row, "二、擦写时间统计", total_columns)
    sequence = "→".join(
        [str(first_value if first_value is not None else 0)]
        + [str(change.after) for change in erasures]
    )
    sheet.cell(row, 1, f"擦写次数：{len(erasures)}　擦写计数值变化：{sequence}").font = Font(
        name="微软雅黑", bold=True, size=10
    )
    row += 1
    headers = ["第几次擦写", "计数值变化", "首次观测时间", "备注 / 依据"]
    for column, header in enumerate(headers, 1):
        _header(sheet, row, column, header)
    if len(erasures) == 0:
        row += 1
        _merge_row(sheet, row, "无跳变", 1, total_columns)
    else:
        for index, change in enumerate(erasures, 1):
            row += 1
            note = (
                "正常单次擦写"
                if change.after - change.before == 1
                else f"异常跳变 +{change.after - change.before}"
            )
            values = [
                f"第{index}次擦写",
                f"{change.before} → {change.after}",
                datetime_label(origin, change.global_second),
                note,
            ]
            for column, value in enumerate(values, 1):
                _body(sheet, row, column, value)
            sheet.merge_cells(
                start_row=row,
                start_column=4,
                end_row=row,
                end_column=max(4, total_columns),
            )

    row += 2
    row = _section(sheet, row, "三、录像时长统计", total_columns)
    if erasures:
        window_headers = []
        for index, change in enumerate(erasures, 1):
            window_headers.append(f"{change.before}→{change.after}\n第{index}次擦写")
        last = erasures[-1]
        window_headers.append(f"{last.before}→{last.after} 之后\n（至今）")
    else:
        window_headers = ["擦写前\n(0 起始·未发生擦写)", "第1次擦写后\n(未发生)"]

    stat_headers = ["通道"] + window_headers + ["当前已录制时长", "全程录像时长", "备注"]
    for column, header in enumerate(stat_headers, 1):
        _header(sheet, row, column, header, height=38)
    row += 1
    for channel in result.config.channels:
        windows = [
            format_duration(value) for value in result.channel_windows[channel]
        ]
        if not erasures:
            windows.append("—（未发生）")
        values = (
            [channel]
            + windows
            + [
                format_duration(result.channel_current[channel]),
                format_duration(result.channel_full[channel]),
                "正常",
            ]
        )
        for column, value in enumerate(values, 1):
            _body(sheet, row, column, value, center=True)
        sheet.cell(row, 1).font = Font(name="微软雅黑", bold=True, size=10)
        row += 1

    note = (
        "「第N次擦写」列表示截至该次擦写发生前的录像时长；"
        "「之后（至今）」表示最后一次擦写后的录像时长。"
        "「当前已录制时长」不包含日志末尾仍未关闭的悬空段。"
    )
    row += 1
    _merge_row(sheet, row, note, 1, total_columns, font=NOTE_FONT)

    row += 2
    row = _section(sheet, row, "四、按天录像明细", total_columns)
    detail_headers = ["日期", "CH1录像时段", "CH1时长", "CH2时长", "CH3时长", "说明"]
    for column, header in enumerate(detail_headers, 1):
        _header(sheet, row, column, header)
    row += 1
    boundaries = tuple(change.global_second for change in erasures)
    for day in result.actual_days:
        day_start = (day - origin).days * DAY
        day_end = day_start + DAY
        segments = {
            channel: _day_window_segments(
                result.timeline.segments,
                channel,
                day_start,
                day_end,
                boundaries,
            )
            for channel in result.config.channels
        }
        if not any(segments.values()):
            continue
        ch1 = segments.get("CH1", [])
        periods = "、".join(
            f"{clock_from_second(start)[:5]}~{clock_from_second(end)[:5]}"
            for start, end in ch1
        )
        dangling = any(
            segment.dangling and segment.end == result.timeline.total_end
            for segment in result.timeline.segments
            if segment.channel == "CH1" and segment.start < day_end and segment.end > day_start
        )
        explanation = "末尾未关闭，按当日结束补算" if dangling else "已完整关闭"
        values = [
            day.strftime("%m-%d"),
            periods or "—",
            format_duration(sum(end - start for start, end in ch1)),
            format_duration(
                sum(end - start for start, end in segments.get("CH2", []))
            ),
            format_duration(
                sum(end - start for start, end in segments.get("CH3", []))
            ),
            explanation,
        ]
        for column, value in enumerate(values, 1):
            _body(sheet, row, column, value, center=column != 2)
        row += 1

    row += 1
    row = _section(sheet, row, "五、异常与说明", total_columns)
    notes: list[tuple[str, Font]] = []
    for anomaly in result.anomalies:
        notes.append((f"{anomaly.kind}：{anomaly.message}", WARN_FONT))
    notes.append(
        (
            f"ACC 断电 {result.acc_off_count} 次；低压断电 "
            f"{result.low_voltage_count} 次。",
            NOTE_FONT,
        )
    )
    notes.extend((text, NOTE_FONT) for text in result.config.notes)
    notes.append(
        (
            "统计口径：录像时长按通道累计；擦写计数只认增长；"
            "断电时开启中的通道按配置的隐式关闭时间结算。",
            NOTE_FONT,
        )
    )
    seen_notes: set[str] = set()
    for text, font in notes:
        if text in seen_notes:
            continue
        seen_notes.add(text)
        _merge_row(sheet, row, f"· {text}", 1, total_columns, font=font)
        sheet.row_dimensions[row].height = 30
        row += 1

    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A1:{get_column_letter(total_columns)}{max(1, row - 1)}"
    sheet.column_dimensions["A"].width = 22
    for column in range(2, 2 + window_count):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    for column in range(2 + window_count, total_columns + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 20


def _section(sheet: Worksheet, row: int, title: str, total_columns: int) -> int:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    cell = sheet.cell(row, 1, title)
    cell.font = Font(name="微软雅黑", bold=True, size=11)
    cell.fill = SECTION_FILL
    cell.alignment = Alignment(vertical="center")
    for column in range(1, total_columns + 1):
        sheet.cell(row, column).fill = SECTION_FILL
        sheet.cell(row, column).border = BORDER
    return row + 1


def _key_value(
    sheet: Worksheet,
    row: int,
    key: str,
    value,
    total_columns: int,
) -> None:
    sheet.cell(row, 1, key).font = Font(name="微软雅黑", bold=True, size=10)
    sheet.cell(row, 1).fill = LIGHT_FILL
    sheet.cell(row, 1).border = BORDER
    sheet.cell(row, 2, value).font = BODY_FONT
    sheet.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.merge_cells(
        start_row=row,
        start_column=2,
        end_row=row,
        end_column=max(2, total_columns),
    )
    length = len(str(value))
    sheet.row_dimensions[row].height = 24 if length < 100 else 42 if length < 200 else 60


def _header(sheet: Worksheet, row: int, column: int, value: str, height: int = 24) -> None:
    cell = sheet.cell(row, column, value)
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = height


def _body(
    sheet: Worksheet,
    row: int,
    column: int,
    value,
    center: bool = False,
) -> None:
    cell = sheet.cell(row, column, value)
    cell.font = BODY_FONT
    cell.border = BORDER
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center",
        wrap_text=True,
    )


def _merge_row(
    sheet: Worksheet,
    row: int,
    value,
    start_column: int,
    end_column: int,
    font: Font = BODY_FONT,
) -> None:
    sheet.merge_cells(
        start_row=row,
        start_column=start_column,
        end_row=row,
        end_column=end_column,
    )
    cell = sheet.cell(row, start_column, value)
    cell.font = font
    cell.alignment = Alignment(wrap_text=True, vertical="center")


def _day_window_segments(
    segments: tuple[Segment, ...],
    channel: str,
    day_start: int,
    day_end: int,
    boundaries: tuple[int, ...],
) -> list[tuple[int, int]]:
    result: list[tuple[int, int]] = []
    for segment in segments:
        if segment.channel != channel:
            continue
        start = max(segment.start, day_start)
        end = min(segment.end, day_end)
        if end <= start:
            continue
        cuts = [start]
        cuts.extend(boundary for boundary in boundaries if start < boundary < end)
        cuts.append(end)
        for part_start, part_end in zip(cuts, cuts[1:]):
            result.append((part_start - day_start, part_end - day_start))
    return result


def _build_progress_workbook(
    results: tuple[DeviceResult, ...],
    config: SiteConfig,
    generated_at: datetime,
) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "统计进度"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:G1")
    title = sheet.cell(1, 1, f"{config.site_name} 统计进度")
    title.font = TITLE_FONT
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26

    latest = max(
        (result.actual_days[-1] for result in results if result.actual_days),
        default=None,
    )
    sheet.merge_cells("A2:G2")
    sheet.cell(
        2,
        1,
        f"最后更新：{generated_at:%Y-%m-%d %H:%M}　　"
        f"数据最新日期：{latest.isoformat() if latest else '—'}",
    ).font = NOTE_FONT

    headers = ["设备", "已统计到", "缺失日期", "擦写次数", "计数值变化", "配置完成时间", "备注"]
    for column, header in enumerate(headers, 1):
        _header(sheet, 3, column, header, height=24)

    row = 4
    for result in results:
        missing = (
            "无"
            if not result.missing_days
            else f"{len(result.missing_days)} 天：" + "、".join(
                day.strftime("%m-%d") for day in result.missing_days
            )
        )
        first = result.timeline.first_value
        sequence = "→".join(
            [str(first if first is not None else 0)]
            + [str(change.after) for change in result.timeline.erasures]
        )
        deltas = [
            change.after - change.before for change in result.timeline.erasures
        ]
        if not deltas:
            note = "配置后尚未出现擦写"
        else:
            maximum = max(deltas)
            note = f"单次最大 +{maximum}" + ("（含异常跳变）" if maximum > 1 else "")
        config_time = "—"
        if result.config_time_label:
            config_time = result.config_time_label[5:16].replace("-", "-")
        values = [
            result.config.display_name,
            result.actual_days[-1].strftime("%m-%d") if result.actual_days else "—",
            missing,
            len(result.timeline.erasures),
            sequence if result.timeline.erasures else f"{first if first is not None else 0}（无跳变）",
            config_time,
            note,
        ]
        for column, value in enumerate(values, 1):
            _body(sheet, row, column, value, center=column not in (1, 3, 5, 7))
        sheet.cell(row, 1).font = Font(name="微软雅黑", bold=True, size=10)
        sheet.row_dimensions[row].height = 24
        row += 1

    row += 1
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
    sheet.cell(
        row,
        1,
        "说明：「当前已录制时长」不含末尾悬空段；缺失日期只登记，不自动补数；"
        "擦写次数按 avgEraseCount 的有效增长次数计算。",
    ).font = NOTE_FONT
    sheet.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row].height = 34

    widths = [18, 13, 30, 11, 20, 18, 28]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "A4"
    sheet.auto_filter.ref = f"A3:G{max(3, row - 2)}"
    workbook.properties.title = f"{config.site_name}统计进度"
    workbook.properties.creator = "device-log-analyzer"
    workbook.properties.created = generated_at
    return workbook

