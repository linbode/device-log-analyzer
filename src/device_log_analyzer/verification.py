from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .metrics import format_duration
from .models import DeviceResult


@dataclass(frozen=True)
class MetricDiff:
    sheet: str
    channel: str
    field: str
    expected: str
    actual: str


def extract_report_metrics(path: Path) -> dict[str, dict[str, dict[str, list[str] | str]]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    extracted: dict[str, dict[str, dict[str, list[str] | str]]] = {}
    for sheet in workbook.worksheets:
        header_row = None
        for row in range(1, min(sheet.max_row, 40) + 1):
            if str(sheet.cell(row, 1).value or "").strip() == "通道":
                header_row = row
                break
        if header_row is None:
            continue
        headers = [
            str(sheet.cell(header_row, column).value or "")
            for column in range(1, sheet.max_column + 1)
        ]
        try:
            current_col = headers.index("当前已录制时长") + 1
            full_col = headers.index("全程录像时长") + 1
        except ValueError:
            continue
        window_values: dict[str, list[str]] = {}
        current_values: dict[str, str] = {}
        full_values: dict[str, str] = {}
        for row in range(header_row + 1, header_row + 8):
            channel = str(sheet.cell(row, 1).value or "").strip()
            if not channel.startswith("CH"):
                continue
            window_values[channel] = [
                str(sheet.cell(row, column).value or "")
                for column in range(2, current_col)
            ]
            current_values[channel] = str(sheet.cell(row, current_col).value or "")
            full_values[channel] = str(sheet.cell(row, full_col).value or "")
        extracted[sheet.title] = {
            channel: {
                "windows": window_values[channel],
                "current": current_values[channel],
                "full": full_values[channel],
            }
            for channel in window_values
        }
    return extracted


def compare_results_to_baseline(
    results: tuple[DeviceResult, ...],
    baseline_path: Path,
) -> list[MetricDiff]:
    expected = extract_report_metrics(baseline_path)
    diffs: list[MetricDiff] = []
    for result in results:
        expected_sheet = expected.get(result.config.display_name)
        if expected_sheet is None:
            diffs.append(
                MetricDiff(
                    result.config.display_name,
                    "—",
                    "工作表",
                    "存在",
                    "缺失",
                )
            )
            continue
        for channel in result.config.channels:
            values = expected_sheet.get(channel)
            if values is None:
                continue
            actual_windows = [
                format_duration(value)
                for value in result.channel_windows.get(channel, ())
            ]
            if not result.timeline.erasures:
                actual_windows.append("—（未发生）")
            expected_windows = list(values["windows"])  # type: ignore[arg-type]
            if len(expected_windows) != len(actual_windows):
                diffs.append(
                    MetricDiff(
                        result.config.display_name,
                        channel,
                        "窗口数量",
                        str(len(expected_windows)),
                        str(len(actual_windows)),
                    )
                )
            for index, (expected_value, actual_value) in enumerate(
                zip(expected_windows, actual_windows), 1
            ):
                if expected_value != actual_value:
                    diffs.append(
                        MetricDiff(
                            result.config.display_name,
                            channel,
                            f"窗口{index}",
                            expected_value,
                            actual_value,
                        )
                    )
            for field, actual_value in (
                ("当前已录制时长", format_duration(result.channel_current[channel])),
                ("全程录像时长", format_duration(result.channel_full[channel])),
            ):
                expected_value = str(values["current" if field.startswith("当前") else "full"])
                if expected_value != actual_value:
                    diffs.append(
                        MetricDiff(
                            result.config.display_name,
                            channel,
                            field,
                            expected_value,
                            actual_value,
                        )
                    )
    return diffs

